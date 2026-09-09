"""
VERIFICATION AGENT ROUTES
Ported from the original Flask app.py -- business logic and SQL are
unchanged; only the routing layer (Flask -> FastAPI-compat) differs.
See app/compat.py for how `request`/`session`/`render_template`/etc. work.
"""
import os
import csv
import json
import base64
import secrets
import traceback
from io import BytesIO, StringIO
from datetime import datetime

import openpyxl

from ..compat import (
    FlaskCompatRouter, request, session, render_template, redirect, url_for,
    flash, jsonify, send_file, send_from_directory,
)
from ..database import get_db, fetchone_dict, fetchall_dict
from ..security import hash_pw, verify_pw, _looks_like_legacy_sha256
from ..config import UPLOAD_FOLDER
from ..helpers import (
    ROLES, ROLE_LABELS, ROLE_RANK, _slugify, get_upload_path,
    generate_company_code, make_company_code, generate_reg_token,
    get_company_credits, consume_credit, log_activity, get_user, get_company,
    get_accessible_ids, require_parent_for_role, classify_period,
    _parse_stmt_date, check_period_duplicate, get_current_quarter,
    save_extraction, get_sa_vcounts, push_notification, get_agent_company_ids,
    _build_verification_excel,
)
from ..auth import login_required, admin_required, superadmin_required, agent_required
from ..extraction import call_ocr_extraction, _compute_file_hash, _compute_content_fingerprint

router = FlaskCompatRouter()

@router.route("/agent/dashboard")
@agent_required
def agent_dashboard():
    """Agent-only dashboard -- assigned companies, verification queue, own stats."""
    agent_id   = session["super_admin_id"]
    agent_name = session.get("super_admin_name", "Agent")
    cids = get_agent_company_ids()

    conn = get_db()
    c = conn.cursor()

    # Companies this agent handles
    companies = []
    if cids:
        ph = ",".join(["%s"] * len(cids))
        c.execute(f"""SELECT co.id, co.name, co.code,
                   COUNT(DISTINCT u.id) as user_count,
                   COUNT(DISTINCT up.id) as upload_count,
                   COUNT(DISTINCT up.id) FILTER (WHERE up.status='done') as done_count,
                   COUNT(DISTINCT mv.id) FILTER (WHERE mv.status='pending') as pending_verif
            FROM companies co
            LEFT JOIN users u ON u.company_id=co.id AND u.status='active'
            LEFT JOIN uploads up ON up.company_id=co.id
            LEFT JOIN manual_verifications mv ON mv.upload_id=up.id
            WHERE co.id IN ({ph}) GROUP BY co.id, co.name, co.code
            ORDER BY co.name""", cids)
        companies = fetchall_dict(c)

    # Agent's own verification stats
    c.execute("""SELECT
        COUNT(*) as total_verified,
        COUNT(*) FILTER (WHERE mv.verified_at::timestamptz::date = CURRENT_DATE) as today,
        COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '7 days') as this_week,
        COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '30 days') as this_month
        FROM manual_verifications mv WHERE mv.verified_by=%s""", (agent_id,))
    my_stats = fetchone_dict(c) or {}

    # Pending verification queue -- only tasks assigned TO this agent
    pending_q = []
    if cids:
        ph = ",".join(["%s"] * len(cids))
        c.execute(f"""SELECT mv.id, mv.status, mv.created_at, mv.excel_downloaded_at,
               u.original_filename, u.stored_filename, u.upload_date,
               co.name as company_name, co.code as company_code,
               d.name as division_name,
               e.stockist_name, e.statement_from_date, e.statement_to_date,
               e.invoice_net, e.doc_type,
               (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) as party_count,
               up_user.full_name as uploader_name
            FROM manual_verifications mv
            JOIN uploads u ON mv.upload_id=u.id
            JOIN companies co ON u.company_id=co.id
            JOIN users up_user ON u.user_id=up_user.id
            LEFT JOIN divisions d ON u.division_id=d.id
            LEFT JOIN extractions e ON e.upload_id=u.id
            WHERE u.company_id IN ({ph})
              AND mv.status='pending'
              AND mv.assigned_to=%s
            ORDER BY mv.created_at ASC""", cids + [agent_id])
        pending_q = fetchall_dict(c)
    else:
        # Superadmin acting as agent -- see all pending assigned to them
        c.execute("""SELECT mv.id, mv.status, mv.created_at, mv.excel_downloaded_at,
               u.original_filename, u.stored_filename, u.upload_date,
               co.name as company_name, co.code as company_code,
               d.name as division_name,
               e.stockist_name, e.statement_from_date, e.statement_to_date,
               e.invoice_net, e.doc_type,
               (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) as party_count,
               up_user.full_name as uploader_name
            FROM manual_verifications mv
            JOIN uploads u ON mv.upload_id=u.id
            JOIN companies co ON u.company_id=co.id
            JOIN users up_user ON u.user_id=up_user.id
            LEFT JOIN divisions d ON u.division_id=d.id
            LEFT JOIN extractions e ON e.upload_id=u.id
            WHERE mv.status='pending' AND mv.assigned_to=%s
            ORDER BY mv.created_at ASC""", (agent_id,))
        pending_q = fetchall_dict(c)

    # Already verified by this agent
    c.execute("""SELECT mv.id, mv.verified_at, mv.corrections_json,
               co.name as company_name, e.stockist_name, e.doc_type,
               e.statement_from_date, e.statement_to_date, e.invoice_net,
               u.original_filename, up_user.full_name as uploader_name
        FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id
        JOIN companies co ON u.company_id=co.id
        JOIN users up_user ON u.user_id=up_user.id
        LEFT JOIN extractions e ON e.upload_id=u.id
        WHERE mv.verified_by=%s
        ORDER BY mv.verified_at DESC LIMIT 20""", (agent_id,))
    my_verified = fetchall_dict(c)

    # Get unique uploaders for filter
    uploaders = []
    if my_verified:
        uploader_set = set()
        for mv in my_verified:
            if mv['uploader_name'] and mv['uploader_name'] not in uploader_set:
                uploader_set.add(mv['uploader_name'])
                uploaders.append(mv['uploader_name'])
        uploaders.sort()

    # Count TRULY unassigned tasks (assigned_to IS NULL only) for sidebar badge
    c.execute("SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL")
    unassigned_count = c.fetchone()[0]

    conn.close()
    return render_template("agent.html", page="dashboard",
        agent_name=agent_name, companies=companies,
        my_stats=my_stats, pending_q=pending_q, my_verified=my_verified,
        unassigned_count=unassigned_count, uploaders=uploaders)


@router.route("/agent/verification-history/export")
@agent_required
def agent_verification_history_export():
    """Export current agent verification history as Excel."""
    agent_id = session["super_admin_id"]
    agent_name = session.get("super_admin_name", "agent")
    conn = get_db()
    c = conn.cursor()
    c.execute("""
        SELECT mv.id as verification_id,
               mv.verified_at,
               mv.status as verification_status,
               u.original_filename as document_name,
               e.stockist_name,
               e.statement_from_date,
               e.statement_to_date,
               e.invoice_net,
               e.doc_type,
               co.name as company_name,
               co.code as company_code,
               up_user.full_name as uploader_name
        FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id = u.id
        JOIN companies co ON u.company_id = co.id
        JOIN users up_user ON u.user_id = up_user.id
        LEFT JOIN extractions e ON e.upload_id = u.id
        WHERE mv.verified_by = %s
        ORDER BY mv.verified_at DESC, mv.id DESC
    """, (agent_id,))
    history_rows = fetchall_dict(c)

    c.execute("""
        SELECT
            COUNT(*) as total_verified,
            COUNT(*) FILTER (WHERE mv.verified_at::timestamptz::date = CURRENT_DATE) as today,
            COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '7 days') as this_week,
            COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '30 days') as this_month
        FROM manual_verifications mv
        WHERE mv.verified_by = %s
    """, (agent_id,))
    agent_summary = fetchone_dict(c) or {}
    conn.close()

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = 'Overview'
    ws1.append(['Agent', agent_name])
    ws1.append(['Total Verified', agent_summary.get('total_verified', 0)])
    ws1.append(['Today', agent_summary.get('today', 0)])
    ws1.append(['This Week', agent_summary.get('this_week', 0)])
    ws1.append(['This Month', agent_summary.get('this_month', 0)])

    ws2 = wb.create_sheet('Verification History')
    ws2.append([
        'Verification ID', 'Verified At', 'Status', 'Company', 'Company Code',
        'Document', 'Stockist', 'Period From', 'Period To', 'Doc Type',
        'Invoice Net', 'Uploader'
    ])
    for row in history_rows:
        ws2.append([
            row['verification_id'], row['verified_at'] or '', row['verification_status'] or '',
            row['company_name'] or '', row['company_code'] or '',
            row['document_name'] or row['stockist_name'] or '',
            row['stockist_name'] or '',
            row['statement_from_date'] or '', row['statement_to_date'] or '',
            row['doc_type'] or '', row['invoice_net'] or 0,
            row['uploader_name'] or ''
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
        as_attachment=True,
        download_name=f"{agent_name.replace(' ', '_').lower()}_verification_history.xlsx",
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@router.route("/agent/verification/<int:mv_id>/download-excel")
@agent_required
def agent_download_excel(mv_id):
    """Agent downloads Excel -- calls shared builder directly (bypasses SA decorator)."""
    conn = get_db()
    c = conn.cursor()
    c.execute("""SELECT mv.upload_id, u.company_id FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id WHERE mv.id=%s""", (mv_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        return "Not found", 404
    return _build_verification_excel(mv_id)


@router.route("/agent/verification/<int:mv_id>/upload-corrections", methods=["POST"])
@agent_required
def agent_upload_corrections(mv_id):
    """Agent uploads corrected Excel -- updates DB, marks verified, redirects to dashboard."""
    cids = get_agent_company_ids()
    conn = get_db()
    c = conn.cursor()
    c.execute("""SELECT mv.upload_id, u.company_id FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id WHERE mv.id=%s""", (mv_id,))
    row = c.fetchone()
    if not row or (cids and row[1] not in cids):
        conn.close()
        return "Access denied", 403
    # Ensure extraction exists before proceeding
    c.execute("SELECT id FROM extractions WHERE upload_id=%s", (row[0],))
    ext = c.fetchone()
    conn.close()
    if not ext:
        flash("No extraction data found for this document. Cannot apply corrections.")
        return redirect(url_for("agent_view_document", mv_id=mv_id))
    # Apply corrections directly (can't call sa route - it has @superadmin_required)
    import json as _json
    from io import BytesIO
    if "excel_file" not in request.files:
        flash("No file uploaded")
        return redirect(url_for("agent_view_document", mv_id=mv_id))
    ef = request.files["excel_file"]
    if not ef.filename.endswith((".xlsx", ".xls")):
        flash("Only .xlsx files accepted")
        return redirect(url_for("agent_view_document", mv_id=mv_id))

    conn2 = get_db(); c2 = conn2.cursor()
    c2.execute("""SELECT mv.*, e.id as ext_id FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id
        LEFT JOIN extractions e ON e.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    mv2 = fetchone_dict(c2)
    if not mv2:
        conn2.close()
        flash("Verification record not found")
        return redirect(url_for("agent_dashboard"))

    try:
        from openpyxl import load_workbook as _lw
        wb2 = _lw(BytesIO(ef.read()), data_only=True)
        ext_id2 = mv2.get("ext_id")
        if not ext_id2:
            c2.execute("SELECT id FROM extractions WHERE upload_id=%s", (mv2["upload_id"],))
            _er = c2.fetchone()
            ext_id2 = _er[0] if _er else None
        if not ext_id2:
            conn2.close()
            flash("⚠ Extraction data not found -- corrections could not be applied.")
            return redirect(url_for("agent_view_document", mv_id=mv_id))

        party_updates = item_updates = 0

        if "Parties" in wb2.sheetnames:
            for row in wb2["Parties"].iter_rows(min_row=2, values_only=True):
                if not row or all(v is None or str(v).strip()=="" for v in row): continue
                pid,name,area,ptype,dl,gst = row[0],str(row[1] or "").strip(),str(row[2] or "").strip(),str(row[3] or "chemist").strip(),str(row[4] or "").strip(),str(row[5] or "").strip()
                if not name: continue
                try: pid_int = int(float(str(pid))) if pid and str(pid).strip() not in ("","0","NEW") else 0
                except: pid_int = 0
                if pid_int > 0:
                    c2.execute("UPDATE parties SET name=%s,area=%s,type=%s,dl_number=%s,gst_number=%s WHERE id=%s",
                        (name,area,ptype,dl,gst,pid_int))
                    party_updates += 1
                else:
                    c2.execute("INSERT INTO parties (extraction_id,name,type,area,dl_number,gst_number,total_quantity,total_amount) VALUES (%s,%s,%s,%s,%s,%s,0,0)",
                        (ext_id2,name,ptype,area,dl,gst))

        c2.execute("SELECT id,name FROM parties WHERE extraction_id=%s", (ext_id2,))
        pname_map = {r[1].strip().upper(): r[0] for r in c2.fetchall()}

        if "Items" in wb2.sheetnames:
            for row in wb2["Items"].iter_rows(min_row=2, values_only=True):
                if not row or all(v is None or str(v).strip()=="" for v in row): continue
                try:
                    iid=row[0]; pid_ref=row[1]; sname=str(row[2] or "").strip()
                    brand=str(row[3] or "").strip(); mfg=str(row[4] or "").strip(); pack=str(row[5] or "").strip()
                    batch=str(row[6] or "").strip(); exp=str(row[7] or "").strip()
                    qty=int(float(str(row[8] or 0))); mrp=float(str(row[9] or 0))
                    rate=float(str(row[10] or 0)); disc=float(str(row[11] or 0)); amt=float(str(row[12] or 0))
                except: continue
                if not brand and qty==0 and amt==0: continue
                try: iid_int = int(float(str(iid))) if iid and str(iid).strip() not in ("","0","NEW") else 0
                except: iid_int = 0
                if iid_int > 0:
                    c2.execute("""UPDATE items SET brand=%s,mfg=%s,pack=%s,batch_no=%s,expiry=%s,
                        quantity=%s,mrp=%s,unit_rate=%s,discount_percent=%s,final_amount=%s WHERE id=%s""",
                        (brand,mfg,pack,batch,exp,qty,mrp,rate,disc,amt,iid_int))
                    item_updates += 1
                else:
                    try: rpid = int(float(str(pid_ref))) if pid_ref else None
                    except: rpid = pname_map.get(str(pid_ref or "").strip().upper()) or pname_map.get(sname.upper())
                    if rpid:
                        c2.execute("INSERT INTO items (party_id,brand,mfg,pack,batch_no,expiry,quantity,mrp,unit_rate,discount_percent,final_amount) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                            (rpid,brand,mfg,pack,batch,exp,qty,mrp,rate,disc,amt))

        # Recalc party totals
        c2.execute("""UPDATE parties SET
            total_quantity=(SELECT COALESCE(SUM(quantity),0) FROM items WHERE party_id=parties.id),
            total_amount=(SELECT COALESCE(SUM(final_amount),0) FROM items WHERE party_id=parties.id)
            WHERE extraction_id=%s""", (ext_id2,))
        # Recalc extraction totals
        c2.execute("""UPDATE extractions SET
            total_quantity=(SELECT COALESCE(SUM(i.quantity),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
            total_amount=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
            invoice_net=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
            net_sale=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id)
            WHERE id=%s""", (ext_id2,))
        # Mark verified
        c2.execute("""UPDATE manual_verifications SET status='verified',
            verified_by=%s, verified_at=%s, excel_uploaded_at=%s
            WHERE id=%s""", (session["super_admin_id"], datetime.now().isoformat(), datetime.now().isoformat(), mv_id))
        c2.execute("UPDATE uploads SET verification_status='verified' WHERE id=%s", (mv2["upload_id"],))
        conn2.commit()
        conn2.close()
        flash(f"✓ Corrections applied -- {party_updates} parties, {item_updates} items updated. Document marked verified.")
    except Exception as _ce:
        try: conn2.rollback(); conn2.close()
        except: pass
        flash(f"⚠ Error applying corrections: {_ce}")
    return redirect(url_for("agent_dashboard"))


@router.route("/agent/verification/<int:mv_id>/inline-save", methods=["POST"])
@agent_required
def agent_inline_save(mv_id):
    """Receive JSON edits from inline editor, apply to DB, mark verified."""
    cids = get_agent_company_ids()
    conn = get_db(); c = conn.cursor()
    c.execute("""SELECT mv.id, mv.upload_id, u.company_id
        FROM manual_verifications mv JOIN uploads u ON mv.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    mv_row = fetchone_dict(c)
    if not mv_row or (cids and mv_row["company_id"] not in cids):
        conn.close(); return jsonify({"error": "Access denied"}), 403

    data = request.get_json(silent=True) or {}
    parties  = data.get("parties", [])   # [{id, name, area, type, dl_number, gst_number}]
    items    = data.get("items", [])     # [{id, party_id, brand, mfg, pack, quantity, unit_rate, discount_percent, final_amount}]
    deleted_party_ids = [int(x) for x in data.get("deleted_party_ids", []) if x]
    deleted_item_ids  = [int(x) for x in data.get("deleted_item_ids",  []) if x]
    new_stockist_name = (data.get("stockist_name", "") or "").strip()
    new_date_from     = (data.get("statement_from_date", "") or "").strip()
    new_date_to       = (data.get("statement_to_date", "") or "").strip()

    try:
        # Delete removed rows (items first, then parties)
        if deleted_item_ids:
            ph = ",".join(["%s"]*len(deleted_item_ids))
            c.execute(f"DELETE FROM items WHERE id IN ({ph})", deleted_item_ids)
        if deleted_party_ids:
            ph = ",".join(["%s"]*len(deleted_party_ids))
            c.execute(f"DELETE FROM items WHERE party_id IN ({ph})", deleted_party_ids)
            c.execute(f"DELETE FROM parties WHERE id IN ({ph})", deleted_party_ids)

        # Get extraction id
        c.execute("SELECT id FROM extractions WHERE upload_id=%s", (mv_row["upload_id"],))
        ext = c.fetchone()
        ext_id = ext[0] if ext else None

        # Upsert parties
        for p in parties:
            pid = int(p.get("id") or 0)
            if pid > 0:
                c.execute("""UPDATE parties SET name=%s, area=%s, type=%s,
                    dl_number=%s, gst_number=%s WHERE id=%s""",
                    (p.get("name",""), p.get("area",""), p.get("type","chemist"),
                     p.get("dl_number",""), p.get("gst_number",""), pid))
            elif ext_id:
                c.execute("""INSERT INTO parties (extraction_id, name, area, type, dl_number, gst_number)
                    VALUES (%s,%s,%s,%s,%s,%s) RETURNING id""",
                    (ext_id, p.get("name",""), p.get("area",""), p.get("type","chemist"),
                     p.get("dl_number",""), p.get("gst_number","")))
                new_pid = c.fetchone()[0]
                # remap temp_id items
                temp_id = p.get("temp_id")
                if temp_id:
                    for item in items:
                        if str(item.get("party_id")) == str(temp_id):
                            item["party_id"] = new_pid

        # Upsert items
        for item in items:
            iid = int(item.get("id") or 0)
            qty   = float(item.get("quantity") or 0)
            rate  = float(item.get("unit_rate") or 0)
            disc  = float(item.get("discount_percent") or 0)
            amt   = float(item.get("final_amount") or 0)
            if iid > 0:
                c.execute("""UPDATE items SET brand=%s, mfg=%s, pack=%s,
                    quantity=%s, unit_rate=%s, discount_percent=%s, final_amount=%s
                    WHERE id=%s""",
                    (item.get("brand",""), item.get("mfg",""), item.get("pack",""),
                     qty, rate, disc, amt, iid))
            else:
                pid_for_item = int(item.get("party_id") or 0)
                if pid_for_item > 0:
                    c.execute("""INSERT INTO items (party_id, brand, mfg, pack,
                        quantity, unit_rate, discount_percent, final_amount)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                        (pid_for_item, item.get("brand",""), item.get("mfg",""),
                         item.get("pack",""), qty, rate, disc, amt))

        # Recalculate party totals
        if ext_id:
            c.execute("SELECT id FROM parties WHERE extraction_id=%s", (ext_id,))
            for (pid,) in c.fetchall():
                c.execute("""SELECT COALESCE(SUM(final_amount),0), COALESCE(SUM(quantity),0)
                    FROM items WHERE party_id=%s""", (pid,))
                p_total, p_qty = c.fetchone()
                c.execute("UPDATE parties SET total_amount=%s, total_quantity=%s WHERE id=%s",
                    (p_total, p_qty, pid))
            # Recalculate extraction total
            c.execute("""SELECT COALESCE(SUM(total_amount),0), COALESCE(SUM(total_quantity),0)
                FROM parties WHERE extraction_id=%s""", (ext_id,))
            e_total, e_qty = c.fetchone()
        # Update stockist name + statement dates if provided
        if new_stockist_name or new_date_from or new_date_to:
            upd_parts = []
            upd_vals = []
            if new_stockist_name: upd_parts.append("stockist_name=%s"); upd_vals.append(new_stockist_name)
            if new_date_from: upd_parts.append("statement_from_date=%s"); upd_vals.append(new_date_from)
            if new_date_to: upd_parts.append("statement_to_date=%s"); upd_vals.append(new_date_to)
            if upd_parts:
                c.execute(f"UPDATE extractions SET {chr(44).join(upd_parts)} WHERE id=%s",
                          upd_vals + [ext_id])
            c.execute("""UPDATE extractions SET total_amount=%s, invoice_net=%s, total_quantity=%s
                WHERE id=%s""", (e_total, e_total, e_qty, ext_id))

        # Mark as verified
        c.execute("""UPDATE manual_verifications SET status='verified',
            verified_at=NOW(), verified_by=%s WHERE id=%s""",
            (session["super_admin_id"], mv_id))
        c.execute("""UPDATE uploads SET verification_status='verified'
            WHERE id=%s""", (mv_row["upload_id"],))
        conn.commit()
        conn.close()
        return jsonify({"success": True, "message": "Saved & marked verified"})
    except Exception as e:
        conn.rollback(); conn.close()
        return jsonify({"error": str(e)}), 500


@router.route("/agent/verification/<int:mv_id>/reextract", methods=["POST"])
@agent_required
def agent_reextract(mv_id):
    """
    Re-run Gemini extraction on the ORIGINAL document from scratch, discarding
    whatever is currently in parties/items for this upload (including any
    manual corrections already made) and replacing it with a fresh AI pass.
    Useful when the first extraction was clearly wrong (e.g. before the
    Pack/Qty prompt fix, or on a document that was rotated/hard to read).
    Resets the verification task back to 'pending' since the data changed.
    """
    cids = get_agent_company_ids()
    conn = get_db(); c = conn.cursor()
    c.execute("""SELECT mv.id, mv.upload_id, u.company_id, u.stored_filename, u.file_type, u.original_filename
        FROM manual_verifications mv JOIN uploads u ON mv.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    mv_row = fetchone_dict(c)
    if not mv_row or (cids and mv_row["company_id"] not in cids):
        conn.close()
        return jsonify({"error": "Access denied"}), 403

    stored_filename = mv_row["stored_filename"]
    fpath = os.path.join(UPLOAD_FOLDER, stored_filename) if stored_filename else None
    if not fpath or not os.path.isfile(fpath):
        conn.close()
        return jsonify({"error": "Original document file is missing from storage -- can't re-extract."}), 404

    try:
        new_data = call_ocr_extraction(fpath, mv_row.get("file_type"), mv_row.get("original_filename") or stored_filename,
                                        company_id=mv_row["company_id"], upload_id=mv_row["upload_id"])
    except Exception as e:
        conn.close()
        return jsonify({"error": f"AI re-extraction failed: {e}"}), 500

    try:
        upload_id = mv_row["upload_id"]
        c.execute("SELECT id FROM extractions WHERE upload_id=%s", (upload_id,))
        ext_row = c.fetchone()
        if ext_row:
            ext_id = ext_row[0]
            c.execute("DELETE FROM items WHERE party_id IN (SELECT id FROM parties WHERE extraction_id=%s)", (ext_id,))
            c.execute("DELETE FROM parties WHERE extraction_id=%s", (ext_id,))
            c.execute("DELETE FROM extractions WHERE id=%s", (ext_id,))
            conn.commit()
        conn.close()

        save_extraction(upload_id, new_data, mv_row["company_id"])

        conn2 = get_db(); c2 = conn2.cursor()
        c2.execute("""UPDATE manual_verifications SET status='pending',
            verified_by=NULL, verified_at=NULL WHERE id=%s""", (mv_id,))
        c2.execute("UPDATE uploads SET verification_status='pending_verification' WHERE id=%s", (upload_id,))
        conn2.commit()
        conn2.close()

        log_activity("reextract_ai", "upload", upload_id,
                      f"Re-extracted with AI by agent (mv #{mv_id})")

        return jsonify({"success": True, "message": "Re-extracted with AI -- please review the new data.", "data": new_data})
    except Exception as e:
        traceback.print_exc()
        try: conn.rollback(); conn.close()
        except Exception: pass
        return jsonify({"error": f"Saved extraction failed: {e}"}), 500


@router.route("/agent/verification/<int:mv_id>/party-items")
@agent_required
def agent_get_party_items(mv_id):
    """Return all parties+items as JSON for inline editor."""
    conn = get_db(); c = conn.cursor()
    c.execute("""SELECT mv.upload_id FROM manual_verifications mv WHERE mv.id=%s""", (mv_id,))
    row = c.fetchone()
    if not row: conn.close(); return jsonify({"error":"Not found"}), 404
    c.execute("SELECT id FROM extractions WHERE upload_id=%s", (row[0],))
    ext = c.fetchone()
    if not ext: conn.close(); return jsonify({"parties":[]})
    ext_id = ext[0]
    c.execute("SELECT * FROM parties WHERE extraction_id=%s ORDER BY id", (ext_id,))
    parties_out = []
    for p in fetchall_dict(c):
        c.execute("SELECT * FROM items WHERE party_id=%s ORDER BY id", (p["id"],))
        parties_out.append({"party": p, "items": fetchall_dict(c)})
    conn.close()
    return jsonify({"parties": parties_out})




@router.route("/agent/verification/<int:mv_id>/reject", methods=["POST"])
@agent_required
def agent_reject_task(mv_id):
    """Agent rejects a verification task with a reason."""
    agent_id = session["super_admin_id"]
    data = request.get_json(silent=True) or {}
    reason = data.get("reason", "").strip()
    reason_label = data.get("reason_label", reason)
    if not reason:
        return jsonify({"error": "Rejection reason required"}), 400
    conn = get_db(); c = conn.cursor()
    c.execute("""SELECT mv.id, mv.upload_id FROM manual_verifications mv WHERE mv.id=%s""", (mv_id,))
    mv_row = fetchone_dict(c)
    if not mv_row:
        conn.close(); return jsonify({"error": "Task not found"}), 404
    c.execute("""UPDATE manual_verifications SET status='rejected', verified_at=NOW(), verified_by=%s
        WHERE id=%s""", (agent_id, mv_id))
    c.execute("""UPDATE uploads SET status='rejected', verification_status='rejected',
        rejection_reason=%s WHERE id=%s""", (reason_label, mv_row["upload_id"]))
    conn.commit(); conn.close()
    return jsonify({"success": True, "message": f"Rejected: {reason_label}"})


@router.route("/agent/profile", methods=["GET", "POST"])
@agent_required
def agent_profile():
    """Agent can change their own password."""
    agent_id = session["super_admin_id"]
    conn = get_db()
    c = conn.cursor()
    if request.method == "POST":
        old_pw  = request.form.get("old_password","").strip()
        new_pw  = request.form.get("new_password","").strip()
        cfm     = request.form.get("confirm_password","").strip()
        c.execute("SELECT password FROM super_admins WHERE id=%s", (agent_id,))
        row = c.fetchone()
        if not row or not verify_pw(row[0], old_pw):
            flash("Current password is incorrect")
        elif new_pw != cfm:
            flash("Passwords do not match")
        elif len(new_pw) < 10:
            flash("Minimum 10 characters required")
        else:
            c.execute("UPDATE super_admins SET password=%s WHERE id=%s",
                      (hash_pw(new_pw), agent_id))
            conn.commit()
            flash("✓ Password updated successfully")
    c.execute("SELECT id, username, full_name, email, role, created_at FROM super_admins WHERE id=%s",
              (agent_id,))
    agent = fetchone_dict(c)
    cids = get_agent_company_ids()
    companies = []
    if cids:
        ph = ",".join(["%s"]*len(cids))
        c.execute(f"SELECT id,name,code FROM companies WHERE id IN ({ph}) ORDER BY name", cids)
        companies = fetchall_dict(c)
    conn.close()
    return render_template("agent.html", page="profile",
        agent_name=session.get("super_admin_name","Agent"),
        agent=agent or {}, companies=companies,
        unassigned_count=unassigned_count_for_page())



@router.route("/agent/verification/<int:mv_id>/claim", methods=["POST"])
@agent_required
def agent_claim_task(mv_id):
    """Agent self-assigns an unassigned or reassigns a pending verification task."""
    agent_id = session["super_admin_id"]
    cids = get_agent_company_ids()
    conn = get_db()
    c = conn.cursor()
    c.execute("""SELECT mv.id, mv.status, mv.assigned_to, u.company_id
        FROM manual_verifications mv JOIN uploads u ON mv.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    row = fetchone_dict(c)
    if not row:
        conn.close()
        return jsonify({"error": "Not found"}), 404
    if cids and row["company_id"] not in cids:
        conn.close()
        return jsonify({"error": "Access denied -- company not in your scope"}), 403
    if row["status"] == "verified":
        conn.close()
        return jsonify({"error": "Already verified"}), 409
    c.execute("UPDATE manual_verifications SET assigned_to=%s WHERE id=%s", (agent_id, mv_id))
    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "Task claimed successfully"})


@router.route("/agent/verification/bulk-claim", methods=["POST"])
@agent_required
def agent_bulk_claim():
    """Agent claims multiple tasks at once."""
    agent_id = session["super_admin_id"]
    data = request.get_json(silent=True) or {}
    mv_ids = data.get("mv_ids", [])
    if not mv_ids:
        return jsonify({"error": "No task IDs provided"}), 400
    conn = get_db(); c = conn.cursor()
    claimed = []; skipped = []
    for mv_id in mv_ids:
        try: mv_id = int(mv_id)
        except: continue
        c.execute("""SELECT mv.id, mv.status, u.company_id FROM manual_verifications mv
            JOIN uploads u ON mv.upload_id=u.id WHERE mv.id=%s AND mv.assigned_to IS NULL""", (mv_id,))
        row = fetchone_dict(c)
        if not row: skipped.append({"id": mv_id, "reason": "not found or already claimed"}); continue
        if row["status"] == "verified": skipped.append({"id": mv_id, "reason": "verified"}); continue
        c.execute("UPDATE manual_verifications SET assigned_to=%s WHERE id=%s", (agent_id, mv_id))
        claimed.append(mv_id)
    conn.commit(); conn.close()
    c2 = get_db().cursor()
    get_db().cursor().execute("SELECT 1")  # dummy to reuse pattern
    conn3 = get_db(); c3 = conn3.cursor()
    c3.execute("SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL")
    new_count = c3.fetchone()[0]; conn3.close()
    return jsonify({"success": True, "claimed": claimed, "skipped": skipped,
                    "claimed_count": len(claimed), "new_unassigned_count": new_count})


@router.route("/agent/verification/unassigned-count")
@agent_required
def agent_unassigned_count_api():
    conn = get_db(); c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL")
    count = c.fetchone()[0]; conn.close()
    return jsonify({"unassigned_count": count})


@router.route("/agent/verification/unassigned")
@agent_required
def agent_unassigned_tasks():
    """List all unassigned/unclaimed tasks the agent can pick up."""
    agent_id = session["super_admin_id"]
    cids = get_agent_company_ids()
    conn = get_db()
    c = conn.cursor()
    agent_id_local = session["super_admin_id"]
    where = "WHERE mv.status='pending' AND mv.assigned_to IS NULL"
    params = []
    if cids:
        ph = ",".join(["%s"] * len(cids))
        where += f" AND u.company_id IN ({ph})"
        params.extend(cids)
    c.execute(f"""SELECT mv.id, mv.assigned_to, mv.created_at,
               u.original_filename, u.upload_date,
               co.name as company_name, co.code as company_code,
               d.name as division_name,
               e.stockist_name, e.statement_from_date, e.statement_to_date,
               e.invoice_net, e.doc_type,
               (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) as party_count,
               up_user.full_name as uploader_name,
               sa.full_name as assigned_agent
        FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id
        JOIN companies co ON u.company_id=co.id
        JOIN users up_user ON u.user_id=up_user.id
        LEFT JOIN divisions d ON u.division_id=d.id
        LEFT JOIN extractions e ON e.upload_id=u.id
        LEFT JOIN super_admins sa ON mv.assigned_to=sa.id
        {where}
        ORDER BY mv.assigned_to ASC NULLS FIRST, mv.created_at ASC""", params)
    tasks = fetchall_dict(c)
    conn.close()
    # Count truly unassigned tasks from DB (not len(tasks) which could be stale)
    conn2 = get_db()
    c2 = conn2.cursor()
    c2.execute("SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL")
    unassigned_count = c2.fetchone()[0]
    conn2.close()
    return render_template("agent.html", page="unassigned",
        agent_name=session.get("super_admin_name", "Agent"),
        tasks=tasks, unassigned_count=unassigned_count)



def unassigned_count_for_page():
    """Quick count helper used by document/profile page renders."""
    conn = get_db(); c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL")
    n = c.fetchone()[0]; conn.close(); return n


@router.route("/agent/document/<int:mv_id>/react")
@agent_required
def agent_view_document_react(mv_id):
    """
    React proof-of-concept for the verification editor -- see the chat
    discussion on incremental React adoption. Deliberately additive: the
    existing Jinja page at /agent/document/<mv_id> is untouched, this is an
    alternate URL so the two can be compared side by side before deciding
    whether to roll React out further. Reuses the exact same JSON endpoints
    (party-items, inline-save, reject, reextract, download-excel, claim) as
    the Jinja version -- no new backend surface for this page.
    """
    conn = get_db()
    c = conn.cursor()
    c.execute("""SELECT mv.id, mv.upload_id, mv.status,
               u.original_filename, u.stored_filename, u.file_type,
               co.name as company_name, co.code as company_code,
               d.name as division_name,
               e.stockist_name, e.statement_from_date, e.statement_to_date,
               e.invoice_net
        FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id
        JOIN companies co ON u.company_id=co.id
        LEFT JOIN divisions d ON u.division_id=d.id
        LEFT JOIN extractions e ON e.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    mv = fetchone_dict(c)
    conn.close()
    if not mv:
        return "Not found", 404
    return render_template("agent_react_poc.html", mv=mv, mv_id=mv_id)


@router.route("/agent/document/<int:mv_id>")
@agent_required
def agent_view_document(mv_id):
    """Agent views the original uploaded document alongside their verification task."""
    conn = get_db()
    c = conn.cursor()
    c.execute("""SELECT mv.id, mv.upload_id, mv.status, mv.excel_downloaded_at,
               u.original_filename, u.stored_filename, u.file_type, u.upload_date,
               co.name as company_name, co.code as company_code,
               d.name as division_name,
               e.stockist_name, e.statement_from_date, e.statement_to_date,
               e.invoice_net, e.doc_type, e.stockist_address, e.stockist_gst,
               (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) as party_count,
               up_user.full_name as uploader_name, up_user.area as uploader_area
        FROM manual_verifications mv
        JOIN uploads u ON mv.upload_id=u.id
        JOIN companies co ON u.company_id=co.id
        JOIN users up_user ON u.user_id=up_user.id
        LEFT JOIN divisions d ON u.division_id=d.id
        LEFT JOIN extractions e ON e.upload_id=u.id
        WHERE mv.id=%s""", (mv_id,))
    mv = fetchone_dict(c)
    if not mv:
        conn.close()
        return "Not found", 404
    # Load extraction + ALL parties with full items (same as extraction_detail)
    c.execute("SELECT * FROM extractions WHERE upload_id=%s", (mv["upload_id"],))
    extraction = fetchone_dict(c)
    parties = []
    if extraction:
        c.execute("SELECT * FROM parties WHERE extraction_id=%s ORDER BY id", (extraction["id"],))
        for p in fetchall_dict(c):
            c.execute("SELECT * FROM items WHERE party_id=%s ORDER BY id", (p["id"],))
            parties.append({"party": p, "items": fetchall_dict(c)})
    conn.close()
    return render_template("agent.html", page="document",
        agent_name=session.get("super_admin_name", "Agent"),
        mv=mv, extraction=extraction, parties=parties, mv_id=mv_id,
        unassigned_count=unassigned_count_for_page())

# ═══════════════════════════════════════════════════════════════════════════════
# COMPANY USER ROUTES
# ═══════════════════════════════════════════════════════════════════════════════
