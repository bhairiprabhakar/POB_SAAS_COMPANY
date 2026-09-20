"""
Statement-document domain (Batch 1b, ported from app/helpers.py).

Business logic behind the legacy statement upload flow, now tenant-scoped:
the tenant database IS the company, so the legacy company_id filters collapse
to the connection itself. Every function takes the caller's tenant connection
(``conn``) and never opens/closes one, matching the saas/ conventions
(see saas/credits.py, saas/audit.py).

Covered here:
  - period helpers (classify_period / _parse_stmt_date / get_current_quarter)
  - save_extraction() -- persist extraction + parties + items, flip the
    upload to 'done'
  - L1 / L2 duplicate checks (file bytes; stockist+period fingerprint)
  - check_period_duplicate() -- overlap-aware period warning
  - verification_counts() -- manual-verification badge counts
"""
import json
from datetime import date, datetime

from .db_utils import fetchall_dict, fetchone_dict

_DATE_FORMATS = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"]

_PERIOD_LABELS = {
    "monthly": "Monthly statement",
    "quarterly": "Quarterly statement",
    "half-yearly": "Half-yearly statement",
    "yearly": "Full-year statement",
}


# ── Period helpers ────────────────────────────────────────────────────────────

def classify_period(from_date_str, to_date_str):
    """Classify a date range into monthly | quarterly | half-yearly | yearly.
    Returns (period_type, days); ("unknown", 0) when either date won't parse."""
    d1 = _parse_stmt_date(from_date_str)
    d2 = _parse_stmt_date(to_date_str)
    if not d1 or not d2:
        return "unknown", 0
    days = (d2 - d1).days
    if days <= 35:
        return "monthly", days
    if days <= 100:
        return "quarterly", days
    if days <= 200:
        return "half-yearly", days
    return "yearly", days


def period_label(period_type, days=None):
    """Friendly label for a period type (used by warnings/UI)."""
    if period_type in _PERIOD_LABELS:
        return _PERIOD_LABELS[period_type]
    return f"Statement ({days or '?'} days)"


def _parse_stmt_date(date_str):
    """Parse a statement date string using the same format list classify_period
    uses, so both functions agree on what a given date string means."""
    if not date_str:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(date_str.strip(), fmt).date()
        except Exception:
            continue
    return None


def get_current_quarter():
    """Indian fiscal year quarters:
      Q1 = April   – June      (months 4–6)
      Q2 = July    – September (months 7–9)
      Q3 = October – December  (months 10–12)
      Q4 = January – March     (months 1–3, belongs to previous FY)
    """
    import calendar
    m = date.today().month
    fy = date.today().year

    if 4 <= m <= 6:
        q, sm, em, fy = 1, 4, 6, fy
    elif 7 <= m <= 9:
        q, sm, em, fy = 2, 7, 9, fy
    elif 10 <= m <= 12:
        q, sm, em, fy = 3, 10, 12, fy
    else:
        q, sm, em, fy = 4, 1, 3, fy   # Jan-Mar = Q4 of previous FY

    q_start = date(fy, sm, 1)
    q_end = date(fy, em, calendar.monthrange(fy, em)[1])
    fy_label = f"FY {fy}-{str(fy + 1)[2:]}" if q < 4 else f"FY {fy - 1}-{str(fy)[2:]}"
    return {
        "quarter": q,
        "label": f"Q{q} {fy_label}",
        "from": q_start.strftime("%Y-%m-%d"),
        "to": q_end.strftime("%Y-%m-%d"),
        "from_display": q_start.strftime("%d/%m/%Y"),
        "to_display": q_end.strftime("%d/%m/%Y"),
    }


# ── Duplicate detection ───────────────────────────────────────────────────────

def check_file_hash_duplicate(conn, file_hash):
    """Level-1 dedup: an earlier upload with the exact same file bytes.
    Tenant-scoped, so 'already uploaded' means within this company. Returns
    the matching upload row (with uploader name) or None."""
    if not file_hash:
        return None
    c = conn.cursor()
    c.execute(
        """SELECT u.id, u.upload_date, us.full_name AS uploader, u.original_filename
           FROM uploads u JOIN users us ON u.user_id=us.id
           WHERE u.file_hash=%s AND u.status='done'
           ORDER BY u.id LIMIT 1""",
        (file_hash,),
    )
    return fetchone_dict(c)


def check_content_fingerprint_duplicate(conn, content_fingerprint):
    """Level-2 dedup: same stockist + same statement period, even under a
    renamed file. Returns the matching upload row (with extraction summary)
    or None. The legacy caller skips this when OCR could not read the
    stockist name (fingerprint of an empty name matches everything)."""
    if not content_fingerprint:
        return None
    c = conn.cursor()
    c.execute(
        """SELECT u.id, u.upload_date, us.full_name AS uploader, u.original_filename,
                  e.stockist_name, e.statement_from_date, e.statement_to_date
           FROM uploads u JOIN users us ON u.user_id=us.id
           JOIN extractions e ON e.upload_id=u.id
           WHERE u.content_fingerprint=%s AND u.status='done'
           ORDER BY u.id LIMIT 1""",
        (content_fingerprint,),
    )
    return fetchone_dict(c)


def check_period_duplicate(conn, stockist_name, from_date, to_date,
                           exclude_upload_id=None):
    """Detect duplicate / overlapping statement periods for the same agency
    (stockist) within this tenant.

    The previous version only flagged an exact from/to string match, which
    missed real duplicates such as:
      - a monthly statement re-uploaded with slightly different OCR'd date
        boundaries (e.g. 1-30 June vs 1-29 June),
      - a monthly statement whose period is already covered by an existing
        quarterly statement for the same agency, or vice versa.

    Matches carry an ``overlap_type``:
      - 'exact'     : identical from/to dates
      - 'contains'  : the new period fully contains an existing one
      - 'contained' : the new period is fully contained within an existing one
      - 'partial'   : the two periods overlap but neither fully contains the other
    """
    new_from = _parse_stmt_date(from_date)
    new_to = _parse_stmt_date(to_date)

    c = conn.cursor()
    q = """SELECT u.id, u.upload_date, u.user_id, us.full_name AS uploader,
                  e.statement_from_date, e.statement_to_date
           FROM extractions e JOIN uploads u ON e.upload_id=u.id
           JOIN users us ON u.user_id=us.id
           WHERE e.stockist_name=%s AND u.status='done'"""
    params = [stockist_name]
    if exclude_upload_id:
        q += " AND u.id!=%s"
        params.append(exclude_upload_id)
    c.execute(q, params)
    rows = fetchall_dict(c)

    if not new_from or not new_to:
        # Can't reliably compare ranges without valid dates on the new upload
        # -- fall back to the previous exact-string-match behaviour.
        return [r for r in rows if r["statement_from_date"] == from_date
                and r["statement_to_date"] == to_date]

    overlaps = []
    for r in rows:
        e_from = _parse_stmt_date(r.get("statement_from_date"))
        e_to = _parse_stmt_date(r.get("statement_to_date"))
        if not e_from or not e_to:
            continue
        if new_to < e_from or new_from > e_to:
            continue  # no overlap
        if new_from == e_from and new_to == e_to:
            r["overlap_type"] = "exact"
        elif new_from <= e_from and new_to >= e_to:
            r["overlap_type"] = "contains"
        elif new_from >= e_from and new_to <= e_to:
            r["overlap_type"] = "contained"
        else:
            r["overlap_type"] = "partial"
        r["existing_period_type"], _ = classify_period(
            r.get("statement_from_date", ""), r.get("statement_to_date", ""))
        overlaps.append(r)

    priority = {"exact": 0, "contains": 1, "contained": 1, "partial": 2}
    overlaps.sort(key=lambda r: priority.get(r["overlap_type"], 3))
    return overlaps


# ── Persistence ───────────────────────────────────────────────────────────────

def save_extraction(conn, upload_id, data):
    """Persist an extraction + its parties + items for ``upload_id`` and flip
    the upload to 'done'. Returns {"parties_saved": n, "items_saved": n}.

    Ported from app/helpers.save_extraction with the company_id dropped: the
    tenant database IS the company. ``data`` is the extraction payload shaped
    by saas/extraction.py (Areas -> parties -> items).
    """
    c = conn.cursor()
    raw_with_debug = json.dumps(data)
    c.execute(
        """INSERT INTO extractions
           (upload_id,stockist_name,stockist_gst,stockist_address,bill_number,bill_date,
            statement_from_date,statement_to_date,total_amount,total_quantity,
            discount_percent,discount_amount,net_sale,sgst,cgst,invoice_net,doc_type,raw_json)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
           RETURNING id""",
        (upload_id,
         str(data.get("stockist_name", "") or "")[:500],
         data.get("stockist_gst", ""),
         str(data.get("stockist_address", "") or "")[:1000],
         data.get("bill_number", ""),
         data.get("bill_date", ""),
         data.get("statement_from_date", ""),
         data.get("statement_to_date", ""),
         float(data.get("total_amount", 0) or 0),
         int(data.get("total_quantity", 0) or 0),
         float(data.get("discount_percent", 0) or 0),
         float(data.get("discount_amount", 0) or 0),
         float(data.get("net_sale", 0) or 0),
         float(data.get("sgst", 0) or 0),
         float(data.get("cgst", 0) or 0),
         float(data.get("invoice_net", 0) or 0),
         data.get("doc_type", "STATEMENT"),
         raw_with_debug),
    )
    ext_id = c.fetchone()[0]

    parties_saved = 0
    items_saved = 0

    for party in data.get("parties", []):
        pname = str(party.get("name", "") or "").strip()
        if not pname:
            continue
        c.execute(
            """INSERT INTO parties
               (extraction_id,name,type,area,dl_number,gst_number,total_quantity,total_amount)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id""",
            (ext_id,
             pname[:500],
             str(party.get("type", "shop") or "shop")[:50],
             str(party.get("area", "") or "")[:200],
             str(party.get("dl_number", "") or "")[:100],
             str(party.get("gst_number", "") or "")[:50],
             int(party.get("party_total_quantity", 0) or 0),
             float(party.get("party_total_amount", 0) or 0)),
        )
        pid = c.fetchone()[0]
        parties_saved += 1

        for item in party.get("items", []):
            brand = str(item.get("brand", "") or "").strip()
            qty = int(item.get("quantity", 0) or 0)
            if not brand and qty == 0:
                continue
            c.execute(
                """INSERT INTO items
                   (party_id,brand,mfg,pack,batch_no,expiry,hsn_code,
                    quantity,mrp,unit_rate,tax_type,discount_percent,final_amount)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (pid,
                 brand[:300],
                 str(item.get("mfg", "") or "")[:200],
                 str(item.get("pack", "") or "")[:100],
                 str(item.get("batch_no", "") or "")[:50],
                 str(item.get("expiry", "") or "")[:30],
                 str(item.get("hsn_code", "") or "")[:20],
                 qty,
                 float(item.get("mrp", 0) or 0),
                 float(item.get("unit_rate", 0) or 0),
                 str(item.get("tax_type", "") or "")[:20],
                 float(item.get("discount_percent", 0) or 0),
                 float(item.get("final_amount", 0) or 0)),
            )
            items_saved += 1

    c.execute("UPDATE uploads SET status='done' WHERE id=%s", (upload_id,))
    conn.commit()
    return {"parties_saved": parties_saved, "items_saved": items_saved}


# ── Verification queue ─────────────────────────────────────────────────────────

def verification_counts(conn) -> dict:
    """Manual-verification counts by status (sidebar/queue badges). Returns an
    empty dict when the tables are unavailable, like the legacy helper."""
    try:
        c = conn.cursor()
        c.execute("SELECT status, COUNT(*) AS cnt FROM manual_verifications GROUP BY status")
        rows = fetchall_dict(c)
        return {r["status"]: r["cnt"] for r in rows}
    except Exception:
        return {}


# ── Verification workflow (Batch 2, ported from app/routers/agent_routes.py) ─
#
# In the merged model verification agents are TENANT USERS with a verification
# role (verifier / verification_agent), not platform super_admins -- the tenant
# DB is the company, and manual_verifications.assigned_to / verified_by
# reference users(id). Every function takes the caller's tenant connection.

_VERIFIER_ROLES = ("verifier", "verification_agent")


def _mv_upload_row(conn, mv_id):
    """Resolve mv -> upload/company-within-tenant. Returns None when missing."""
    c = conn.cursor()
    c.execute(
        "SELECT mv.id, mv.upload_id, mv.status, mv.assigned_to, mv.division_id "
        "FROM manual_verifications mv WHERE mv.id=%s",
        (mv_id,),
    )
    return fetchone_dict(c)


def ensure_manual_verification(conn, upload_id: int, division_id=None) -> int:
    """Create a pending manual-verification row for ``upload_id`` if it does
    not exist yet, then auto-assign it to the least-loaded active verifier
    (tenant user with a verification role). Returns the mv id."""
    c = conn.cursor()
    c.execute("SELECT id FROM manual_verifications WHERE upload_id=%s", (upload_id,))
    row = c.fetchone()
    mv_id = row[0] if row else None
    if not mv_id:
        c.execute(
            "INSERT INTO manual_verifications (upload_id, division_id, status) "
            "VALUES (%s, %s, 'pending') RETURNING id",
            (upload_id, division_id),
        )
        mv_id = c.fetchone()[0]
    agent_id = find_least_loaded_agent(conn)
    if agent_id:
        c.execute(
            "UPDATE manual_verifications SET assigned_to=%s "
            "WHERE id=%s AND status='pending' AND assigned_to IS NULL",
            (agent_id, mv_id),
        )
    c.execute("UPDATE uploads SET verification_status='pending_verification' WHERE id=%s", (upload_id,))
    conn.commit()
    return mv_id


def find_least_loaded_agent(conn) -> int | None:
    """Least-loaded active tenant user holding a verification role, or None."""
    c = conn.cursor()
    c.execute(
        """SELECT u.id
           FROM users u
           LEFT JOIN roles r ON r.id = u.role_id
           LEFT JOIN manual_verifications mv2 ON mv2.assigned_to=u.id AND mv2.status='pending'
           WHERE u.status='active' AND r.name IN ('verifier','verification_agent')
           GROUP BY u.id
           ORDER BY COUNT(mv2.id) FILTER (WHERE mv2.id IS NOT NULL) ASC, u.id ASC
           LIMIT 1"""
    )
    row = c.fetchone()
    return row[0] if row else None


def verification_detail(conn, mv_id: int) -> dict | None:
    """Full verification-task payload for the editor page -- mv + upload meta +
    extraction + parties with items. Ported from legacy agent_view_document."""
    c = conn.cursor()
    c.execute(
        """SELECT mv.id, mv.upload_id, mv.status, mv.excel_downloaded_at,
                  u.original_filename, u.stored_filename, u.file_type, u.upload_date,
                  d.name AS division_name,
                  e.stockist_name, e.statement_from_date, e.statement_to_date,
                  e.invoice_net, e.doc_type, e.stockist_address, e.stockist_gst,
                  (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) AS party_count,
                  up_user.full_name AS uploader_name, up_user.area AS uploader_area
           FROM manual_verifications mv
           JOIN uploads u ON mv.upload_id=u.id
           JOIN users up_user ON u.user_id=up_user.id
           LEFT JOIN divisions d ON u.division_id=d.id
           LEFT JOIN extractions e ON e.upload_id=u.id
           WHERE mv.id=%s""",
        (mv_id,),
    )
    mv = fetchone_dict(c)
    if not mv:
        return None
    c.execute("SELECT * FROM extractions WHERE upload_id=%s", (mv["upload_id"],))
    extraction = fetchone_dict(c)
    parties = []
    if extraction:
        c.execute("SELECT * FROM parties WHERE extraction_id=%s ORDER BY id", (extraction["id"],))
        for p in fetchall_dict(c):
            c.execute("SELECT * FROM items WHERE party_id=%s ORDER BY id", (p["id"],))
            parties.append({"party": p, "items": fetchall_dict(c)})
    return {"mv": mv, "extraction": extraction, "parties": parties}


def verification_queue(conn, agent_id: int) -> list:
    """Pending verification tasks assigned to this agent (dashboard queue)."""
    c = conn.cursor()
    c.execute(
        """SELECT mv.id, mv.status, mv.created_at, mv.excel_downloaded_at,
                  u.original_filename, u.upload_date,
                  d.name AS division_name,
                  e.stockist_name, e.statement_from_date, e.statement_to_date,
                  e.invoice_net, e.doc_type,
                  (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) AS party_count,
                  up_user.full_name AS uploader_name
           FROM manual_verifications mv
           JOIN uploads u ON mv.upload_id=u.id
           JOIN users up_user ON u.user_id=up_user.id
           LEFT JOIN divisions d ON u.division_id=d.id
           LEFT JOIN extractions e ON e.upload_id=u.id
           WHERE mv.status='pending' AND mv.assigned_to=%s
           ORDER BY mv.created_at ASC""",
        (agent_id,),
    )
    return fetchall_dict(c)


def unassigned_tasks(conn) -> list:
    """Pending verification tasks no agent has claimed yet."""
    c = conn.cursor()
    c.execute(
        """SELECT mv.id, mv.assigned_to, mv.created_at,
                  u.original_filename, u.upload_date,
                  d.name AS division_name,
                  e.stockist_name, e.statement_from_date, e.statement_to_date,
                  e.invoice_net, e.doc_type,
                  (SELECT COUNT(*) FROM parties p WHERE p.extraction_id=e.id) AS party_count,
                  up_user.full_name AS uploader_name
           FROM manual_verifications mv
           JOIN uploads u ON mv.upload_id=u.id
           JOIN users up_user ON u.user_id=up_user.id
           LEFT JOIN divisions d ON u.division_id=d.id
           LEFT JOIN extractions e ON e.upload_id=u.id
           WHERE mv.status='pending' AND mv.assigned_to IS NULL
           ORDER BY mv.created_at ASC""",
    )
    return fetchall_dict(c)


def unassigned_count(conn) -> int:
    c = conn.cursor()
    c.execute(
        "SELECT COUNT(*) FROM manual_verifications WHERE status='pending' AND assigned_to IS NULL"
    )
    return c.fetchone()[0] or 0


def agent_stats(conn, agent_id: int) -> dict:
    """Agent's own verification stats (today / week / month / total)."""
    c = conn.cursor()
    c.execute(
        """SELECT COUNT(*) AS total_verified,
                  COUNT(*) FILTER (WHERE mv.verified_at::timestamptz::date = CURRENT_DATE) AS today,
                  COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '7 days') AS this_week,
                  COUNT(*) FILTER (WHERE mv.verified_at::timestamptz >= NOW() - INTERVAL '30 days') AS this_month
           FROM manual_verifications mv WHERE mv.verified_by=%s""",
        (agent_id,),
    )
    return fetchone_dict(c) or {}


def verification_history(conn, agent_id: int) -> list:
    c = conn.cursor()
    c.execute(
        """SELECT mv.id AS verification_id,
                  mv.verified_at,
                  mv.status AS verification_status,
                  u.original_filename AS document_name,
                  e.stockist_name,
                  e.statement_from_date,
                  e.statement_to_date,
                  e.invoice_net,
                  e.doc_type,
                  up_user.full_name AS uploader_name
           FROM manual_verifications mv
           JOIN uploads u ON mv.upload_id = u.id
           JOIN users up_user ON u.user_id = up_user.id
           LEFT JOIN extractions e ON e.upload_id = u.id
           WHERE mv.verified_by = %s
           ORDER BY mv.verified_at DESC, mv.id DESC""",
        (agent_id,),
    )
    return fetchall_dict(c)


def claim_task(conn, mv_id: int, agent_id: int):
    """Return (ok: bool, error: str|None, status: int). Ported from
    agent_claim_task minus the cross-company scope check (tenant DB IS the
    company)."""
    row = _mv_upload_row(conn, mv_id)
    if not row:
        return False, "Not found", 404
    if row["status"] == "verified":
        return False, "Already verified", 409
    c = conn.cursor()
    c.execute("UPDATE manual_verifications SET assigned_to=%s WHERE id=%s", (agent_id, mv_id))
    conn.commit()
    return True, None, 200


def bulk_claim(conn, mv_ids, agent_id: int) -> dict:
    claimed = []
    skipped = []
    c = conn.cursor()
    for raw in mv_ids or []:
        try:
            mv_id = int(raw)
        except (TypeError, ValueError):
            continue
        c.execute(
            "SELECT id, status FROM manual_verifications WHERE id=%s AND assigned_to IS NULL",
            (mv_id,),
        )
        row = c.fetchone()
        if not row:
            skipped.append({"id": mv_id, "reason": "not found or already claimed"})
            continue
        if row[1] == "verified":
            skipped.append({"id": mv_id, "reason": "verified"})
            continue
        c.execute("UPDATE manual_verifications SET assigned_to=%s WHERE id=%s", (agent_id, mv_id))
        claimed.append(mv_id)
    conn.commit()
    return {
        "claimed": claimed,
        "skipped": skipped,
        "claimed_count": len(claimed),
        "new_unassigned_count": unassigned_count(conn),
    }


def reject_verification(conn, mv_id: int, agent_id: int, reason: str, reason_label: str = ""):
    """Port of legacy agent_reject_task: the task is marked rejected and the
    backing upload is moved to rejected with the reason."""
    row = _mv_upload_row(conn, mv_id)
    if not row:
        return False, "Task not found", 404
    label = reason_label or reason
    c = conn.cursor()
    c.execute(
        "UPDATE manual_verifications SET status='rejected', verified_at=%s, verified_by=%s WHERE id=%s",
        (datetime.now().isoformat(), agent_id, mv_id),
    )
    c.execute(
        "UPDATE uploads SET status='rejected', verification_status='rejected', "
        "rejection_reason=%s WHERE id=%s",
        (label, row["upload_id"]),
    )
    conn.commit()
    return True, None, 200


def save_inline_corrections(conn, mv_id: int, agent_id: int, data: dict):
    """Apply inline-editor JSON edits from the verification page and mark the
    task verified. Ported from legacy agent_inline_save (no company guard).
    Returns (ok, error_message|None)."""
    parties = data.get("parties", []) or []
    items = data.get("items", []) or []
    deleted_party_ids = [int(x) for x in (data.get("deleted_party_ids", []) or []) if x]
    deleted_item_ids = [int(x) for x in (data.get("deleted_item_ids", []) or []) if x]
    new_stockist_name = (data.get("stockist_name", "") or "").strip()
    new_date_from = (data.get("statement_from_date", "") or "").strip()
    new_date_to = (data.get("statement_to_date", "") or "").strip()

    row = _mv_upload_row(conn, mv_id)
    if not row:
        return False, "Verification task not found"

    c = conn.cursor()
    if deleted_item_ids:
        c.execute("DELETE FROM items WHERE id = ANY(%s)", (deleted_item_ids,))
    if deleted_party_ids:
        c.execute("DELETE FROM items WHERE party_id = ANY(%s)", (deleted_party_ids,))
        c.execute("DELETE FROM parties WHERE id = ANY(%s)", (deleted_party_ids,))

    c.execute("SELECT id FROM extractions WHERE upload_id=%s", (row["upload_id"],))
    ext = c.fetchone()
    ext_id = ext[0] if ext else None

    for p in parties:
        pid = int(p.get("id") or 0)
        if pid > 0:
            c.execute(
                "UPDATE parties SET name=%s, area=%s, type=%s, dl_number=%s, gst_number=%s WHERE id=%s",
                (p.get("name", "") or "", p.get("area", "") or "", p.get("type", "chemist") or "chemist",
                 p.get("dl_number", "") or "", p.get("gst_number", "") or "", pid),
            )
        elif ext_id:
            c.execute(
                "INSERT INTO parties (extraction_id, name, area, type, dl_number, gst_number) "
                "VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",
                (ext_id, p.get("name", "") or "", p.get("area", "") or "",
                 p.get("type", "chemist") or "chemist", p.get("dl_number", "") or "",
                 p.get("gst_number", "") or ""),
            )
            new_pid = c.fetchone()[0]
            temp_id = p.get("temp_id")
            if temp_id:
                for item in items:
                    if str(item.get("party_id")) == str(temp_id):
                        item["party_id"] = new_pid

    for item in items:
        iid = int(item.get("id") or 0)
        qty = float(item.get("quantity") or 0)
        rate = float(item.get("unit_rate") or 0)
        disc = float(item.get("discount_percent") or 0)
        amt = float(item.get("final_amount") or 0)
        if iid > 0:
            c.execute(
                "UPDATE items SET brand=%s, mfg=%s, pack=%s, quantity=%s, unit_rate=%s, "
                "discount_percent=%s, final_amount=%s WHERE id=%s",
                (item.get("brand", "") or "", item.get("mfg", "") or "", item.get("pack", "") or "",
                 qty, rate, disc, amt, iid),
            )
        else:
            pid_for_item = int(item.get("party_id") or 0)
            if pid_for_item > 0 and ext_id:
                c.execute(
                    "INSERT INTO items (party_id, brand, mfg, pack, quantity, unit_rate, "
                    "discount_percent, final_amount) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
                    (pid_for_item, item.get("brand", "") or "", item.get("mfg", "") or "",
                     item.get("pack", "") or "", qty, rate, disc, amt),
                )

    e_total = e_qty = 0
    if ext_id:
        c.execute("SELECT id FROM parties WHERE extraction_id=%s", (ext_id,))
        for (pid,) in c.fetchall():
            c.execute(
                "SELECT COALESCE(SUM(final_amount),0), COALESCE(SUM(quantity),0) FROM items WHERE party_id=%s",
                (pid,),
            )
            p_total, p_qty = c.fetchone()
            c.execute("UPDATE parties SET total_amount=%s, total_quantity=%s WHERE id=%s",
                      (p_total, p_qty, pid))
        c.execute(
            "SELECT COALESCE(SUM(total_amount),0), COALESCE(SUM(total_quantity),0) "
            "FROM parties WHERE extraction_id=%s",
            (ext_id,),
        )
        e_total, e_qty = c.fetchone()
        if new_stockist_name or new_date_from or new_date_to:
            upd_parts, upd_vals = [], []
            if new_stockist_name:
                upd_parts.append("stockist_name=%s")
                upd_vals.append(new_stockist_name)
            if new_date_from:
                upd_parts.append("statement_from_date=%s")
                upd_vals.append(new_date_from)
            if new_date_to:
                upd_parts.append("statement_to_date=%s")
                upd_vals.append(new_date_to)
            if upd_parts:
                c.execute("UPDATE extractions SET " + ", ".join(upd_parts) + " WHERE id=%s",
                          upd_vals + [ext_id])
        c.execute(
            "UPDATE extractions SET total_amount=%s, invoice_net=%s, total_quantity=%s WHERE id=%s",
            (e_total, e_total, e_qty, ext_id),
        )

    c.execute(
        "UPDATE manual_verifications SET status='verified', verified_at=%s, verified_by=%s WHERE id=%s",
        (datetime.now().isoformat(), agent_id, mv_id),
    )
    c.execute("UPDATE uploads SET verification_status='verified' WHERE id=%s", (row["upload_id"],))
    conn.commit()
    return True, None


def apply_excel_corrections(conn, mv_id: int, agent_id: int, wb):
    """Apply a corrected workbook (Parties / Items sheets) and mark the task
    verified. Ported from legacy agent_upload_corrections (no company guard).
    Returns (party_updates, item_updates) or raises on failure."""
    c = conn.cursor()
    mv = fetchone_from(conn, "manual_verifications", mv_id)
    if not mv:
        raise ValueError("Verification record not found")
    ext_id = None
    c.execute("SELECT id FROM extractions WHERE upload_id=%s", (mv["upload_id"],))
    er = c.fetchone()
    ext_id = er[0] if er else None
    if not ext_id:
        raise ValueError("Extraction data not found -- corrections could not be applied")

    party_updates = item_updates = 0
    if "Parties" in wb.sheetnames:
        for row in wb["Parties"].iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            pid = row[0]
            name = str(row[1] or "").strip()
            area = str(row[2] or "").strip()
            ptype = str(row[3] or "chemist").strip()
            dl = str(row[4] or "").strip()
            gst = str(row[5] or "").strip()
            if not name:
                continue
            try:
                pid_int = int(float(str(pid))) if pid and str(pid).strip() not in ("", "0", "NEW") else 0
            except Exception:
                pid_int = 0
            if pid_int > 0:
                c.execute(
                    "UPDATE parties SET name=%s, area=%s, type=%s, dl_number=%s, gst_number=%s WHERE id=%s",
                    (name, area, ptype, dl, gst, pid_int),
                )
                party_updates += 1
            else:
                c.execute(
                    "INSERT INTO parties (extraction_id, name, type, area, dl_number, gst_number, "
                    "total_quantity, total_amount) VALUES (%s,%s,%s,%s,%s,%s,0,0)",
                    (ext_id, name, ptype, area, dl, gst),
                )

    c.execute("SELECT id, name FROM parties WHERE extraction_id=%s", (ext_id,))
    pname_map = {r[1].strip().upper(): r[0] for r in c.fetchall()}

    if "Items" in wb.sheetnames:
        for row in wb["Items"].iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            try:
                iid, pid_ref, sname = row[0], row[1], str(row[2] or "").strip()
                brand = str(row[3] or "").strip()
                mfg = str(row[4] or "").strip()
                pack = str(row[5] or "").strip()
                batch = str(row[6] or "").strip()
                exp = str(row[7] or "").strip()
                qty = int(float(str(row[8] or 0)))
                mrp = float(str(row[9] or 0))
                rate = float(str(row[10] or 0))
                disc = float(str(row[11] or 0))
                amt = float(str(row[12] or 0))
            except Exception:
                continue
            if not brand and qty == 0 and amt == 0:
                continue
            try:
                iid_int = int(float(str(iid))) if iid and str(iid).strip() not in ("", "0", "NEW") else 0
            except Exception:
                iid_int = 0
            if iid_int > 0:
                c.execute(
                    "UPDATE items SET brand=%s, mfg=%s, pack=%s, batch_no=%s, expiry=%s, "
                    "quantity=%s, mrp=%s, unit_rate=%s, discount_percent=%s, final_amount=%s WHERE id=%s",
                    (brand, mfg, pack, batch, exp, qty, mrp, rate, disc, amt, iid_int),
                )
                item_updates += 1
            else:
                try:
                    rpid = int(float(str(pid_ref))) if pid_ref else None
                except Exception:
                    rpid = pname_map.get(str(pid_ref or "").strip().upper()) or pname_map.get(sname.upper())
                if rpid:
                    c.execute(
                        "INSERT INTO items (party_id, brand, mfg, pack, batch_no, expiry, quantity, "
                        "mrp, unit_rate, discount_percent, final_amount) "
                        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                        (rpid, brand, mfg, pack, batch, exp, qty, mrp, rate, disc, amt),
                    )

    c.execute(
        """UPDATE parties SET
               total_quantity=(SELECT COALESCE(SUM(quantity),0) FROM items WHERE party_id=parties.id),
               total_amount=(SELECT COALESCE(SUM(final_amount),0) FROM items WHERE party_id=parties.id)
           WHERE extraction_id=%s""",
        (ext_id,),
    )
    c.execute(
        """UPDATE extractions SET
               total_quantity=(SELECT COALESCE(SUM(i.quantity),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
               total_amount=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
               invoice_net=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id),
               net_sale=(SELECT COALESCE(SUM(i.final_amount),0) FROM items i JOIN parties p ON i.party_id=p.id WHERE p.extraction_id=extractions.id)
           WHERE id=%s""",
        (ext_id,),
    )
    now = datetime.now().isoformat()
    c.execute(
        "UPDATE manual_verifications SET status='verified', verified_by=%s, verified_at=%s, "
        "excel_uploaded_at=%s WHERE id=%s",
        (agent_id, now, now, mv_id),
    )
    c.execute("UPDATE uploads SET verification_status='verified' WHERE id=%s", (mv["upload_id"],))
    conn.commit()
    return party_updates, item_updates


def reextract_discard(conn, upload_id: int):
    """Wipe parties/items/extraction for an upload (used before re-running
    OCR). Returns True if an extraction existed and was removed."""
    c = conn.cursor()
    c.execute("SELECT id FROM extractions WHERE upload_id=%s", (upload_id,))
    ext = c.fetchone()
    if not ext:
        return False
    c.execute("DELETE FROM items WHERE party_id IN (SELECT id FROM parties WHERE extraction_id=%s)", (ext[0],))
    c.execute("DELETE FROM parties WHERE extraction_id=%s", (ext[0],))
    c.execute("DELETE FROM extractions WHERE id=%s", (ext[0],))
    conn.commit()
    return True


def reset_verification_pending(conn, mv_id: int, upload_id: int):
    c = conn.cursor()
    c.execute(
        "UPDATE manual_verifications SET status='pending', verified_by=NULL, verified_at=NULL WHERE id=%s",
        (mv_id,),
    )
    c.execute("UPDATE uploads SET verification_status='pending_verification' WHERE id=%s", (upload_id,))
    conn.commit()


def documents(conn, accessible, date_from="", date_to="", status_filter="all"):
    """Tenant document list with status counts. ``accessible`` is the caller's
    visible user-id list from saas/scoping (None = unrestricted). Ported from
    the legacy /documents route query shapes."""
    clause = ""
    params = []
    if accessible is not None:
        clause += " AND u.user_id = ANY(%s)"
        params.append(list(accessible))
    if date_from:
        clause += " AND DATE(u.upload_date) >= %s"
        params.append(date_from)
    if date_to:
        clause += " AND DATE(u.upload_date) <= %s"
        params.append(date_to)
    if status_filter and status_filter not in ("all", ""):
        if status_filter == "verified_stage":
            clause += " AND u.verification_status='verified'"
        elif status_filter == "pending_verif":
            clause += " AND u.verification_status='pending_verification'"
        else:
            clause += " AND u.status=%s"
            params.append(status_filter)

    c = conn.cursor()
    c.execute(
        """SELECT u.*, us.full_name AS uploader_name, r.name AS uploader_role,
                  e.stockist_name, e.invoice_net,
                  e.statement_from_date, e.statement_to_date,
                  rej.full_name AS rejected_by_name,
                  mv.verified_at
           FROM uploads u
           JOIN users us ON u.user_id=us.id
           LEFT JOIN roles r ON r.id=us.role_id
           LEFT JOIN extractions e ON e.upload_id=u.id
           LEFT JOIN users rej ON u.rejected_by=rej.id
           LEFT JOIN manual_verifications mv ON mv.upload_id=u.id
           WHERE 1=1 """ + clause + """
           ORDER BY u.upload_date DESC LIMIT 200""",
        params,
    )
    uploads = fetchall_dict(c)

    count_clause = " WHERE u.user_id = ANY(%s)" if accessible is not None else ""
    count_params = [list(accessible)] if accessible is not None else []
    c.execute(
        "SELECT u.status, u.verification_status, COUNT(*) AS cnt FROM uploads u" + count_clause
        + " GROUP BY u.status, u.verification_status",
        count_params,
    )
    status_counts = {}
    for r in fetchall_dict(c):
        status_counts[r["status"]] = status_counts.get(r["status"], 0) + r["cnt"]
        if r.get("verification_status") == "verified":
            status_counts["verified_stage"] = status_counts.get("verified_stage", 0) + r["cnt"]
        if r.get("verification_status") == "pending_verification":
            status_counts["pending_verif"] = status_counts.get("pending_verif", 0) + r["cnt"]
    return uploads, status_counts


def extraction_payload(conn, up_id: int) -> dict | None:
    """Upload + extraction + parties+items + debug_info for the extraction
    detail page. Returns None when the upload does not exist."""
    c = conn.cursor()
    c.execute(
        """SELECT u.*, us.full_name AS uploader_name, us.area, us.region,
                  r.name AS uploader_role
           FROM uploads u JOIN users us ON u.user_id=us.id
           LEFT JOIN roles r ON r.id=us.role_id
           WHERE u.id=%s""",
        (up_id,),
    )
    upload = fetchone_dict(c)
    if not upload:
        return None
    c.execute("SELECT * FROM extractions WHERE upload_id=%s", (up_id,))
    ext = fetchone_dict(c)
    parties = []
    if ext:
        c.execute("SELECT * FROM parties WHERE extraction_id=%s ORDER BY id", (ext["id"],))
        for p in fetchall_dict(c):
            c.execute("SELECT * FROM items WHERE party_id=%s ORDER BY id", (p["id"],))
            parties.append({"party": p, "items": fetchall_dict(c)})
    debug_info = {}
    if ext and ext.get("raw_json"):
        try:
            debug_info = (json.loads(ext["raw_json"]) or {}).get("_debug", {})
        except Exception:
            pass
    return {"upload": upload, "extraction": ext, "parties": parties, "debug_info": debug_info}


def fetchone_from(conn, table: str, row_id: int) -> dict | None:
    c = conn.cursor()
    c.execute(f"SELECT * FROM {table} WHERE id=%s", (row_id,))
    return fetchone_dict(c)


def build_verification_workbook(conn, mv_id: int, company_code: str = "tenant"):
    """Build the verification Excel workbook (Meta / Parties / Items /
    Instructions), ported from app/helpers._build_verification_excel with the
    companies join dropped (the tenant DB IS the company). Marks the task's
    excel_downloaded_at. Returns (BytesIO, filename)."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
    from openpyxl.utils import get_column_letter

    c = conn.cursor()
    c.execute(
        """SELECT mv.*, u.original_filename, u.stored_filename,
                  e.id AS ext_id, e.stockist_name, e.stockist_gst, e.stockist_address,
                  e.statement_from_date, e.statement_to_date,
                  e.invoice_net, e.total_quantity,
                  e.bill_number, e.bill_date
           FROM manual_verifications mv
           JOIN uploads u ON mv.upload_id=u.id
           LEFT JOIN extractions e ON e.upload_id=u.id
           WHERE mv.id=%s""",
        (mv_id,),
    )
    mv = fetchone_dict(c)
    if not mv:
        raise LookupError("Verification record not found")

    c.execute(
        """SELECT p.*, array_agg(row_to_json(i)) AS items_json
           FROM parties p
           LEFT JOIN items i ON i.party_id=p.id
           WHERE p.extraction_id=%s
           GROUP BY p.id ORDER BY p.name""",
        (mv["ext_id"],),
    )
    parties_raw = fetchall_dict(c)

    c.execute("UPDATE manual_verifications SET excel_downloaded_at=%s WHERE id=%s",
              (datetime.now().isoformat(), mv_id))
    conn.commit()

    wb = Workbook()
    thin = Side(style="thin", color="CCCCCC")
    B = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr(ws, row, cols, bg="1A3560"):
        for i, t in enumerate(cols, 1):
            cl = ws.cell(row, i, t)
            cl.fill = PatternFill("solid", fgColor=bg)
            cl.font = Font(bold=True, color="FFFFFF", size=10)
            cl.alignment = Alignment(horizontal="center", vertical="center")
            cl.border = B

    def cell(ws, row, col, val, editable=False):
        cl = ws.cell(row, col, val)
        cl.border = B
        cl.font = Font(size=9, color="000000" if editable else "444444")
        if editable:
            cl.fill = PatternFill("solid", fgColor="FFFDE7")
        if isinstance(val, float):
            cl.number_format = "#,##0.00"
        cl.alignment = Alignment(vertical="center")
        return cl

    ws_meta = wb.active
    ws_meta.title = "Meta"
    ws_meta["A1"] = "MANUAL VERIFICATION SHEET"
    ws_meta["A1"].font = Font(bold=True, size=14, color="1A3560")
    ws_meta["A1"].alignment = Alignment(horizontal="center")
    ws_meta.merge_cells("A1:D1")
    meta_rows = [
        ("Verification ID", mv_id),
        ("Company", company_code),
        ("Stockist / Agency", mv["stockist_name"] or ""),
        ("GST", mv["stockist_gst"] or ""),
        ("Address", mv["stockist_address"] or ""),
        ("Statement From", mv["statement_from_date"] or ""),
        ("Statement To", mv["statement_to_date"] or ""),
        ("Bill Number", mv["bill_number"] or ""),
        ("Original File", mv["original_filename"] or ""),
        ("Invoice Net (₹)", mv["invoice_net"] or 0),
        ("Total Qty", mv["total_quantity"] or 0),
        ("Instructions", "Yellow cells are editable. Do NOT change column headers or mv_id."),
    ]
    for r, (lbl, val) in enumerate(meta_rows, 3):
        ws_meta.cell(r, 1, lbl).font = Font(bold=True, size=10)
        ws_meta.cell(r, 2, val).font = Font(size=10)
    ws_meta.column_dimensions["A"].width = 22
    ws_meta.column_dimensions["B"].width = 50

    ws_p = wb.create_sheet("Parties")
    hdr(ws_p, 1, ["party_id", "Store Name", "Location / Area", "Type", "DL Number",
                  "GST Number", "Total Qty", "Total Amount ₹"])
    for r, p in enumerate(parties_raw, 2):
        cell(ws_p, r, 1, p["id"])
        cell(ws_p, r, 2, p["name"] or "", editable=True)
        cell(ws_p, r, 3, p["area"] or "", editable=True)
        cell(ws_p, r, 4, p["type"] or "chemist", editable=True)
        cell(ws_p, r, 5, p["dl_number"] or "", editable=True)
        cell(ws_p, r, 6, p["gst_number"] or "", editable=True)
        cell(ws_p, r, 7, p["total_quantity"] or 0)
        cell(ws_p, r, 8, float(p["total_amount"] or 0))
    for col in ws_p.columns:
        ws_p.column_dimensions[get_column_letter(col[0].column)].width = 22

    ws_i = wb.create_sheet("Items")
    hdr(ws_i, 1, ["item_id", "party_id", "Store Name", "Brand / Product",
                  "Company (MFG)", "Pack", "Batch No", "Expiry",
                  "Qty", "MRP ₹", "Rate ₹", "Disc%", "Amount ₹"])
    row = 2
    for p in parties_raw:
        items = []
        try:
            raw = p.get("items_json") or []
            if isinstance(raw, str):
                raw = json.loads(raw)
            for itm in raw:
                if isinstance(itm, str):
                    itm = json.loads(itm)
                if itm:
                    items.append(itm)
        except Exception:
            pass
        for itm in items:
            cell(ws_i, row, 1, itm.get("id", ""))
            cell(ws_i, row, 2, p["id"])
            cell(ws_i, row, 3, p["name"] or "")
            cell(ws_i, row, 4, itm.get("brand", ""), editable=True)
            cell(ws_i, row, 5, itm.get("mfg", ""), editable=True)
            cell(ws_i, row, 6, itm.get("pack", ""), editable=True)
            cell(ws_i, row, 7, itm.get("batch_no", ""), editable=True)
            cell(ws_i, row, 8, itm.get("expiry", ""), editable=True)
            cell(ws_i, row, 9, itm.get("quantity", 0), editable=True)
            cell(ws_i, row, 10, float(itm.get("mrp", 0) or 0), editable=True)
            cell(ws_i, row, 11, float(itm.get("unit_rate", 0) or 0), editable=True)
            cell(ws_i, row, 12, float(itm.get("discount_percent", 0) or 0), editable=True)
            cell(ws_i, row, 13, float(itm.get("final_amount", 0) or 0), editable=True)
            row += 1
    for col in ws_i.columns:
        ws_i.column_dimensions[get_column_letter(col[0].column)].width = 18

    ws_inst = wb.create_sheet("Instructions")
    instructions = [
        ("HOW TO USE THIS SHEET", ""),
        ("", ""),
        ("1. Yellow cells", "Are editable -- correct any wrong values here"),
        ("2. ID columns (col A)", "Do NOT edit existing IDs. For NEW rows leave ID blank or put 0"),
        ("3. Add a new STORE", "New row in Parties sheet -- ID=0, fill Store Name and Area"),
        ("4. Add a new BRAND", "New row in Items sheet -- ID=0, party_id = the store ID from Parties col A"),
        ("5. Remove a row", "Do NOT delete rows -- set Qty=0 and Amount=0 instead"),
        ("6. Do NOT", "Rename sheet tabs or change column headers"),
        ("7. After editing", "Save the file and upload via Verification > Upload Corrections"),
        ("8. Verification ID", f"Must remain: {mv_id} (shown in Meta sheet)"),
    ]
    for r, (lbl, val) in enumerate(instructions, 1):
        ws_inst.cell(r, 1, lbl).font = Font(bold=True, size=11 if r == 1 else 10)
        ws_inst.cell(r, 2, val).font = Font(size=10)
    ws_inst.column_dimensions["A"].width = 22
    ws_inst.column_dimensions["B"].width = 55

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"verify_{company_code}_{mv_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return buf, fname


def verification_history_workbook(conn, agent_id: int, agent_name: str):
    """Agent verification-history export workbook (ported from legacy
    agent_verification_history_export; the company columns are dropped
    because the tenant DB is one company). Returns (BytesIO, filename)."""
    from io import BytesIO
    from openpyxl import Workbook

    history_rows = verification_history(conn, agent_id)
    summary = agent_stats(conn, agent_id)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Overview"
    ws1.append(["Agent", agent_name])
    ws1.append(["Total Verified", summary.get("total_verified", 0)])
    ws1.append(["Today", summary.get("today", 0)])
    ws1.append(["This Week", summary.get("this_week", 0)])
    ws1.append(["This Month", summary.get("this_month", 0)])

    ws2 = wb.create_sheet("Verification History")
    ws2.append(["Verification ID", "Verified At", "Status", "Document",
                "Stockist", "Period From", "Period To", "Doc Type",
                "Invoice Net", "Uploader"])
    for row in history_rows:
        ws2.append([
            row["verification_id"], row["verified_at"] or "", row["verification_status"] or "",
            row["document_name"] or row["stockist_name"] or "",
            row["stockist_name"] or "",
            row["statement_from_date"] or "", row["statement_to_date"] or "",
            row["doc_type"] or "", row["invoice_net"] or 0,
            row["uploader_name"] or "",
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    fname = f"{str(agent_name or 'agent').replace(' ', '_').lower()}_verification_history.xlsx"
    return output, fname